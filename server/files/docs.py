"""Reading a file as text, and writing text back into it - for plain files and Word documents.

A Word document is shown as one line per paragraph, with headings marked `#`, `##`... like
Markdown, because that is a form both you and a model can edit. Writing it back changes only the
paragraphs whose text changed: the rest of the document - its other paragraphs, tables, images,
headers, styles - is the original file, untouched. A changed paragraph keeps the formatting of its
first run; a new one takes the style of the paragraph before it.
"""

from __future__ import annotations

import difflib
import hashlib
import re
from pathlib import Path
from typing import Any

from server.models.files import FileKind

MAX_BYTES = 2_000_000
HEADING = re.compile(r"^(#{1,6}) (.*)$")
DOCX_NOTE = (
    "Word document: one line per paragraph, headings marked with #. Tables, images and "
    "comments are kept as they are, but are not shown here and cannot be edited."
)


class Unreadable(ValueError):
    """Said to the person: why this file cannot be opened as text."""


def version(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()[:16]


def kind_of(path: Path) -> FileKind:
    return "docx" if path.suffix.lower() == ".docx" else "text"


def read(path: Path) -> tuple[FileKind, str, bytes]:
    data = path.read_bytes()
    if len(data) > MAX_BYTES:
        raise Unreadable(f"{path.name} is {len(data) / 1e6:.1f}MB; Files opens files up to 2MB.")
    if kind_of(path) == "docx":
        return "docx", "\n".join(_line(p) for p in _document(path).paragraphs), data
    if b"\x00" in data[:8000]:
        raise Unreadable(f"{path.name} is a binary file, not text.")
    try:
        return "text", data.decode("utf-8"), data
    except UnicodeDecodeError:
        return "text", data.decode("latin-1"), data


def write(path: Path, kind: FileKind, old_text: str, new_text: str) -> None:
    if kind == "docx":
        _write_docx(path, old_text.split("\n"), new_text.split("\n"))
    else:
        path.write_text(new_text, encoding="utf-8", newline="")


def _document(path: Path) -> Any:
    import docx

    return docx.Document(str(path))


def _level(paragraph: Any) -> int:
    name = str(getattr(paragraph.style, "name", "") or "")
    if name.startswith("Heading "):
        digits = name.removeprefix("Heading ").strip()
        return int(digits) if digits.isdigit() else 0
    return 1 if name == "Title" else 0


def _line(paragraph: Any) -> str:
    level = _level(paragraph)
    return f"{'#' * level} {paragraph.text}" if level else paragraph.text


def _split(line: str) -> tuple[int, str]:
    match = HEADING.match(line)
    return (len(match.group(1)), match.group(2)) if match else (0, line)


def _set_text(paragraph: Any, text: str) -> None:
    """New words, the first run's formatting: bold stays bold, the font stays the font."""
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(text)
        return
    runs[0].text = text
    for run in runs[1:]:
        run._element.getparent().remove(run._element)


def _restyle(paragraph: Any, level: int) -> None:
    if level and _level(paragraph) != level:
        paragraph.style = f"Heading {level}"
    elif not level and _level(paragraph):
        paragraph.style = "Normal"


def _insert_after(anchor: Any, document: Any) -> Any:
    from docx.oxml import OxmlElement
    from docx.text.paragraph import Paragraph

    element = OxmlElement("w:p")
    if anchor is None:
        document.element.body.insert(0, element)
        return Paragraph(element, document._body)
    anchor._p.addnext(element)
    created = Paragraph(element, anchor._parent)
    created.style = anchor.style
    return created


def _write_docx(path: Path, old: list[str], new: list[str]) -> None:
    document = _document(path)
    paragraphs = list(document.paragraphs)
    if [_line(p) for p in paragraphs] != old:
        raise Unreadable(f"{path.name} changed since it was read; open it again.")
    matcher = difflib.SequenceMatcher(None, old, new, autojunk=False)
    for tag, i1, i2, j1, j2 in reversed(matcher.get_opcodes()):  # back to front keeps indexes valid
        if tag == "equal":
            continue
        targets, lines = paragraphs[i1:i2], new[j1:j2]
        for paragraph, line in zip(targets, lines, strict=False):
            level, text = _split(line)
            _set_text(paragraph, text)
            _restyle(paragraph, level)
        for paragraph in targets[len(lines) :]:
            paragraph._element.getparent().remove(paragraph._element)
        anchor = targets[-1] if targets else (paragraphs[i1 - 1] if i1 else None)
        for line in lines[len(targets) :]:
            level, text = _split(line)
            anchor = _insert_after(anchor, document)
            _set_text(anchor, text)
            _restyle(anchor, level)
    document.save(str(path))
